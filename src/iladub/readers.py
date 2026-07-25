"""Document readers: turn a source file into text for knowledge-guided extraction.

The reference implementation reads plain text natively. Richer formats
(pdf/docx/xlsx/html) are supported when their optional dependencies are
installed (``pip install 'iladub[readers]'``); otherwise a clear error is
raised rather than silently degrading.
"""
from __future__ import annotations

import csv as _csv
import os

TEXT_SUFFIXES = {".txt", ".md", ".text", ""}


def read_document(path: str) -> str:
    """Return the text content of ``path``, dispatching on file extension."""
    suffix = os.path.splitext(path)[1].lower()
    if suffix in TEXT_SUFFIXES:
        with open(path, "r", encoding="utf-8") as fh:
            return fh.read()
    if suffix in (".html", ".htm"):
        return _read_html(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    # Unknown: best-effort as UTF-8 text.
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()


def _read_html(path: str) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("Reading HTML needs beautifulsoup4: pip install 'iladub[readers]'") from exc
    with open(path, "r", encoding="utf-8") as fh:
        return BeautifulSoup(fh.read(), "html.parser").get_text(" ", strip=True)


def _read_pdf(path: str) -> str:  # pragma: no cover - optional dependency
    try:
        from pdfminer.high_level import extract_text
    except ImportError as exc:
        raise RuntimeError("Reading PDF needs pdfminer.six: pip install 'iladub[readers]'") from exc
    return extract_text(path)


def _read_docx(path: str) -> str:  # pragma: no cover - optional dependency
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("Reading DOCX needs python-docx: pip install 'iladub[readers]'") from exc
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def _read_xlsx(path: str) -> str:  # pragma: no cover - optional dependency
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Reading XLSX needs openpyxl: pip install 'iladub[readers]'") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(rows)


def read_csv_surface_concepts(path: str) -> list:
    """Format adapter: a CSV's header row names the concepts; each data cell is a value.
    Returns region-anchored SurfaceConcepts. Deterministic — no model calls. This is the
    format-coupled boundary; the grounding portal downstream is format-agnostic."""
    from .ground import SurfaceConcept

    out: list[SurfaceConcept] = []
    with open(path, newline="", encoding="utf-8") as fh:
        reader = _csv.DictReader(fh)
        for r, row in enumerate(reader, start=1):
            for header, cell in row.items():
                out.append(SurfaceConcept(text=header, value=cell,
                                          region="row%d:col-%s" % (r, header)))
    return out
